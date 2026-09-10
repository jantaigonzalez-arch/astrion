#!/usr/bin/env python3
"""
EL XLS DEL SAT → LOS CSV QUE ESPERA `npm run sat:catalogos`.

    python3 scripts/sat-xls-a-csv.py catCFDI.xls  --salida ~/sat-csv

── POR QUÉ ESTE PASO EXISTE ────────────────────────────────────────────────

El SAT publica los catálogos del Anexo 20 en UN archivo `.xls` de formato
antiguo (OLE2, unos 48 MB) con veintiocho hojas dentro. El cargador de la
aplicación lee CSV, y con razón: leer `.xls` exigiría una dependencia más en el
candado —el mismo candado que ya tumbó un despliegue— para un trámite que se
hace dos veces al año.

Así que la conversión vive aquí, fuera de la aplicación, y es reproducible.

── DOS COSAS QUE EL SAT HACE Y HAY QUE DESHACER ───────────────────────────

1. LOS CATÁLOGOS GRANDES VIENEN PARTIDOS. Los códigos postales en dos hojas
   (`c_CodigoPostal_Parte_1` y `_2`) y las colonias en TRES. Se concatenan, y el
   preámbulo de las partes siguientes se descarta — si no, el título de la
   hoja 2 entraría como si fuera un código postal.

2. EL PREÁMBULO. Cada hoja abre con su título y la versión del catálogo antes
   de los encabezados. Se busca la fila que contiene la columna clave en vez de
   suponer que es la primera.

La cabecera de dos filas —«Aplica para tipo persona» arriba, «Física / Moral»
abajo— NO se toca aquí: se copia tal cual y la resuelve el cargador, que es
donde vive la decisión de qué columna es cuál.

── REQUISITOS ──────────────────────────────────────────────────────────────

    soffice --headless --convert-to xlsx catCFDI.xls     (LibreOffice)
    python3 -m pip install openpyxl                       (si no está)

El script hace la conversión solo si le das el `.xls` y encuentra `soffice`.
"""
import csv
import os
import shutil
import subprocess
import sys
from pathlib import Path

# Qué hoja alimenta qué archivo. Las listas se concatenan en orden.
CATALOGOS = {
    "c_RegimenFiscal": ["c_RegimenFiscal"],
    "c_UsoCFDI": ["c_UsoCFDI"],
    "c_FormaPago": ["c_FormaPago"],
    "c_MetodoPago": ["c_MetodoPago"],
    "c_Pais": ["c_Pais"],
    "c_CodigoPostal": ["c_CodigoPostal_Parte_1", "c_CodigoPostal_Parte_2"],
    "c_Colonia": ["C_Colonia_1", "C_Colonia_2", "C_Colonia_3"],
}

# La columna que identifica el arranque de los datos en cada hoja.
CLAVE = {
    "c_RegimenFiscal": "c_RegimenFiscal",
    "c_UsoCFDI": "c_UsoCFDI",
    "c_FormaPago": "c_FormaPago",
    "c_MetodoPago": "c_MetodoPago",
    "c_Pais": "c_Pais",
    "c_CodigoPostal": "c_CodigoPostal",
    "c_Colonia": "c_Colonia",
}


def norm(s):
    """Sin acentos ni puntuación: para casar un encabezado con su nombre."""
    import unicodedata

    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return "".join(c for c in s.lower() if c.isalnum())


def celda(v):
    """Una celda como texto, sin la hora que Excel le cuelga a las fechas."""
    if v is None:
        return ""
    t = str(v)
    # `2022-01-01 00:00:00` → `2022-01-01`. El cargador acepta ISO tal cual.
    if t.endswith(" 00:00:00"):
        return t[:-9]
    return t


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)

    origen = Path(sys.argv[1]).expanduser()
    salida = Path(
        sys.argv[sys.argv.index("--salida") + 1] if "--salida" in sys.argv else "sat-csv"
    ).expanduser()
    salida.mkdir(parents=True, exist_ok=True)

    # `.xls` antiguo: se convierte con LibreOffice, que es lo que hay en un Mac.
    if origen.suffix.lower() == ".xls":
        if not shutil.which("soffice"):
            sys.exit("Falta LibreOffice (`soffice`) para convertir el .xls a .xlsx")
        print(f"▸ Convirtiendo {origen.name} a xlsx (tarda: son ~48 MB)…")
        perfil = salida / ".lo-profile"
        subprocess.run(
            ["soffice", "--headless", f"-env:UserInstallation=file://{perfil}",
             "--convert-to", "xlsx", "--outdir", str(salida), str(origen)],
            check=True, capture_output=True,
        )
        origen = salida / (origen.stem + ".xlsx")

    import openpyxl

    wb = openpyxl.load_workbook(origen, read_only=True)
    faltan = []

    for destino, hojas in CATALOGOS.items():
        presentes = [h for h in hojas if h in wb.sheetnames]
        if not presentes:
            faltan.append(destino)
            continue

        filas_totales = 0
        with open(salida / f"{destino}.csv", "w", newline="", encoding="utf8") as f:
            w = csv.writer(f)
            cabecera_escrita = False

            for hoja in presentes:
                ws = wb[hoja]
                clave = norm(CLAVE[destino])
                encontrado = False

                for fila in ws.iter_rows(values_only=True):
                    valores = [celda(c) for c in fila]

                    if not encontrado:
                        # El preámbulo se descarta hasta dar con la clave.
                        if any(norm(v) == clave for v in valores):
                            encontrado = True
                            # La cabecera se escribe UNA vez, la de la primera
                            # parte. Las siguientes repiten la suya y sería una
                            # fila de datos con nombres de columna dentro.
                            if not cabecera_escrita:
                                w.writerow(valores)
                                cabecera_escrita = True
                                # La sub-cabecera —primera celda vacía— también
                                # viaja: la resuelve el cargador.
                                continue
                            else:
                                continue
                        continue

                    if not any(v.strip() for v in valores):
                        continue
                    w.writerow(valores)
                    filas_totales += 1

                if not encontrado:
                    faltan.append(f"{destino} (hoja {hoja} sin la columna {CLAVE[destino]})")

        print(f"  {destino:<20} {filas_totales:>8} filas   ({', '.join(presentes)})")

    if faltan:
        print("\n⚠ No se pudo con:")
        for f in faltan:
            print(f"    {f}")

    print(f"\n▸ CSV en {salida}")
    print("  Ahora:  npm run sat:catalogos -- --dir " + str(salida))


if __name__ == "__main__":
    main()
